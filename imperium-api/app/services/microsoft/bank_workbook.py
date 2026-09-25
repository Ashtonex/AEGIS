"""Bank & Project Money workbook - a two-way Excel mirror in the Data Room.

AEGIS builds one workbook from the bank statement, its splits / cash uses,
per-project money and the books check, and keeps it in the Financial Data
Room's SharePoint folder (the same connected site/drive data_room_sync.py
uses, so no new Graph permission is needed). Pinned as a Teams tab, people
see the live figures - and can edit:

    Bank Lines           Category, Project, Who, Note
    Splits & Cash Uses   Project, Category, Amount, Date, What it paid for on
                         each part (clear Amount to remove it), plus new parts
                         in the spare rows at the bottom (Line = the bank
                         Line ID or its first 8 characters)

Split / cash-use edits go through bank_books.replace_allocations per line (a
savepoint each), so they follow exactly the rules of the AEGIS split editor.
New rows applied in a cycle whose republish didn't land are remembered in
applied_new_row_keys so a re-read can't add them twice.

Everything else is locked (protected sheets; Excel for the web respects
that). Category and Project are dropdowns fed from a hidden-ish Lists sheet.

sync() is the single entry point (after-change job, 2-minute cron, nightly,
"Publish now"). One cycle:

1. Take the per-organisation lease so two cycles never overlap.
2. Ask SharePoint for the file's eTag. If it differs from what AEGIS last
   uploaded, download it and compare each editable cell with the snapshot of
   what AEGIS last published:
     cell == snapshot          -> nobody touched it
     cell == AEGIS now         -> already in AEGIS
     AEGIS changed it too      -> conflict: AEGIS wins, row flagged
     otherwise                 -> apply through the normal tagging path
                                  (reconciliation.tag_lines + bank_books.sync)
   Every outcome is logged in finance.bank_workbook_changes.
3. Rebuild the workbook (with a Sync status column) and upload it with
   If-Match on the eTag it read. If someone edited in those few seconds the
   upload is refused (412) and nothing is overwritten - the next cycle reads
   their edit and tries again.
4. Store the new eTag + snapshot, release the lease.

Edits are credited to the AEGIS user whose email matches the file's
"last modified by" (Microsoft only reports the last editor of the whole file,
not per cell); otherwise they are recorded with no user and that name.
"""

from __future__ import annotations

import hashlib
import io
import json
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import Any, Optional
from urllib.parse import quote
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from core.logging import logger
from app.services.finance import bank_books
from app.services.finance import bank_reconciliation as reconciliation
from app.services.finance.general_ledger import GeneralLedgerError
from app.services.microsoft import data_room_sync, sharepoint
from app.services.microsoft.errors import GraphConflictError, GraphNotFoundError
from app.services.microsoft.graph_client import GraphClient

WORKBOOK_FILE_NAME = "AEGIS Bank & Project Money.xlsx"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = '#,##0.00;[Red]-#,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EDIT_HEADER_FILL = PatternFill("solid", fgColor="B45309")
EDIT_CELL_FILL = PatternFill("solid", fgColor="FEF3C7")
LEASE_MINUTES = 5
STATUS_WINDOW_DAYS = 30

CATEGORY_LABELS = {
    "client_receipt": "Client receipt", "capital_injection": "Capital / owner injection",
    "internal_transfer": "Internal transfer", "cash_withdrawal": "Cash withdrawal",
    "supplier_payment": "Supplier / materials", "subcontractor": "Subcontractor",
    "equipment_hire": "Equipment / plant hire", "fuel_transport": "Fuel & transport",
    "salaries_wages": "Salaries & wages", "tax_statutory": "Tax & statutory",
    "bank_charges": "Bank charges & IMTT", "card_purchase": "Card / POS purchase",
    "owner_drawings": "Owner drawings", "tithe_donation": "Tithe / donation", "refund": "Refund",
    "reversal": "Bank reversal", "site_petty_cash": "Handed to site petty cash", "other": "Other",
}
# What a person can pick for a whole bank line (site petty cash is only a use of cash)
EDITABLE_CATEGORIES = [c for c in CATEGORY_LABELS if c != "site_petty_cash"]
EDITABLE_FIELDS = ("category", "project", "who", "note")
FIELD_TO_COLUMN = {"category": "category", "project": "project_id", "who": "counterparty_name", "note": "notes"}
FIELD_LIMITS = {"who": 200, "note": 2000}

BANK_LINE_HEADERS = ["Line ID", "Date", "Reference", "Description", "Money in", "Money out", "Balance",
                     "Category", "Project", "Who", "Note", "Split into parts", "Journal", "Bank account", "Sync status"]
EDITABLE_HEADERS = {"Category": "category", "Project": "project", "Who": "who", "Note": "note"}

# Splits & Cash Uses: existing parts can be changed / cleared, new ones typed
# into spare rows (a protected sheet can't grow its table, so the spare rows
# are part of it). "Line" takes a bank Line ID, or its first 8+ characters.
SPLIT_HEADERS = ["Allocation ID", "Line", "Line date", "Reference", "Line amount", "Kind", "Project", "Category",
                 "Amount", "Date", "What it paid for", "Sync status"]
SPLIT_EDITABLE = {"Project": "project", "Category": "category", "Amount": "amount", "Date": "date",
                  "What it paid for": "description"}
SPARE_SPLIT_ROWS = 100
SHORT_ID_MIN = 8
# A use of withdrawn cash can't itself be one of these
NOT_A_CASH_USE = ("cash_withdrawal", "client_receipt", "capital_injection", "reversal", "internal_transfer")


def _label(category: Optional[str]) -> str:
    return CATEGORY_LABELS.get(category or "", (category or "").replace("_", " "))


def _num(value: Any) -> Optional[float]:
    return None if value is None else float(value)


def _clean(value: Any) -> Optional[str]:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------

async def collect(db: AsyncSession, org_id: str) -> dict:
    q = lambda sql, **p: db.execute(text(sql), {"org_id": org_id, **p})
    lines = [dict(r) for r in (await q(r"""
        SELECT l.id, l.transaction_date, l.reference,
               regexp_replace(COALESCE(l.description, ''), '\s+', ' ', 'g') AS description,
               l.counterparty_name, l.amount, l.category, l.project_id, p.name AS project, l.notes,
               SUM(l.amount) OVER (PARTITION BY l.import_id ORDER BY l.transaction_date, l.line_number) AS balance,
               je.journal_number,
               (SELECT count(*) FROM finance.bank_line_allocations a WHERE a.line_id = l.id) AS parts,
               ca.account_name AS bank_account,
               (SELECT CASE c.outcome
                          WHEN 'applied' THEN 'Applied from Excel ' || to_char(c.created_at AT TIME ZONE :tz, 'DD Mon HH24:MI')
                          WHEN 'rejected' THEN 'Not applied: ' || c.message
                          ELSE 'Conflict: ' || c.message END
                FROM finance.bank_workbook_changes c
                WHERE c.line_id = l.id AND c.field <> 'split' AND c.created_at > NOW() - make_interval(days => :window)
                ORDER BY c.created_at DESC LIMIT 1) AS sync_status
        FROM finance.bank_statement_lines l
        LEFT JOIN projects.projects p ON p.id = l.project_id
        LEFT JOIN finance.journal_entries je ON je.id = l.gl_journal_id
        LEFT JOIN finance.cash_accounts ca ON ca.id = l.cash_account_id
        WHERE l.organization_id = :org_id
        ORDER BY l.transaction_date, l.line_number
    """, window=STATUS_WINDOW_DAYS, tz=await _timezone(db, org_id))).mappings()]
    tz = await _timezone(db, org_id)
    allocations = [dict(r) for r in (await q("""
        SELECT a.id, a.line_id, l.transaction_date AS line_date, l.reference, l.amount AS line_amount,
               l.category AS line_category, a.project_id, p.name AS project, a.category, a.amount,
               a.allocation_date, a.description,
               (SELECT CASE c.outcome
                          WHEN 'applied' THEN 'Applied from Excel ' || to_char(c.created_at AT TIME ZONE :tz, 'DD Mon HH24:MI')
                          WHEN 'rejected' THEN 'Not applied: ' || c.message
                          ELSE 'Conflict: ' || c.message END
                FROM finance.bank_workbook_changes c
                WHERE c.allocation_id = a.id AND c.created_at > NOW() - make_interval(days => :window)
                ORDER BY c.created_at DESC LIMIT 1) AS sync_status
        FROM finance.bank_line_allocations a
        JOIN finance.bank_statement_lines l ON l.id = a.line_id
        LEFT JOIN projects.projects p ON p.id = a.project_id
        WHERE a.organization_id = :org_id
        ORDER BY l.transaction_date, a.created_at
    """, tz=tz, window=STATUS_WINDOW_DAYS)).mappings()]
    # New split rows typed in Excel that couldn't be applied - they don't
    # survive a republish, so list them on Read Me for re-entry.
    problems = [dict(r) for r in (await q("""
        SELECT to_char(created_at AT TIME ZONE :tz, 'DD Mon HH24:MI') AS at, new_value, message
        FROM finance.bank_workbook_changes
        WHERE organization_id = :org_id AND field = 'split' AND allocation_id IS NULL AND outcome = 'rejected'
          AND created_at > NOW() - INTERVAL '7 days'
        ORDER BY created_at DESC LIMIT 20
    """, tz=tz)).mappings()]
    projects = [dict(r) for r in (await q("""
        SELECT p.id, p.name, p.project_code, p.client_name, p.status, COALESCE(p.contract_value, 0) AS contract_value,
               COALESCE((SELECT sum(certified_amount) FROM finance.progress_claims pc
                         WHERE pc.project_id = p.id AND pc.status IN ('certified', 'paid') AND NOT pc.is_deleted), 0) AS certified,
               COALESCE((SELECT sum(net_claim_amount) FROM finance.progress_claims pc
                         WHERE pc.project_id = p.id AND pc.status = 'paid' AND NOT pc.is_deleted), 0) AS collected,
               COALESCE((SELECT sum(amount) FROM finance.cost_transactions ct WHERE ct.project_id = p.id), 0) AS actual_cost,
               COALESCE((SELECT sum(amount) FROM finance.bank_statement_lines l WHERE l.project_id = p.id AND l.amount > 0), 0) AS bank_in,
               COALESCE((SELECT -sum(amount) FROM finance.bank_statement_lines l WHERE l.project_id = p.id AND l.amount < 0), 0) AS bank_out,
               COALESCE((SELECT sum(a.amount) FROM finance.bank_line_allocations a
                         JOIN finance.bank_statement_lines l ON l.id = a.line_id AND l.category = 'cash_withdrawal'
                         WHERE a.project_id = p.id), 0) AS cash_used
        FROM projects.projects p
        WHERE p.organization_id = :org_id AND p.is_deleted = false
        ORDER BY collected DESC, p.name
    """)).mappings()]
    audit = await bank_books.audit(db, org_id=org_id)
    return {"lines": lines, "allocations": allocations, "projects": projects, "audit": audit, "problems": problems}


async def _timezone(db: AsyncSession, org_id: str) -> str:
    tz = (await db.execute(
        text("""
            SELECT oi.calendar_timezone FROM core.organisation_integrations oi
            WHERE oi.organization_id = :o AND oi.provider = 'microsoft365'
              AND EXISTS (SELECT 1 FROM pg_timezone_names t WHERE t.name = oi.calendar_timezone)
        """),
        {"o": org_id},
    )).scalar()
    return tz or "Africa/Harare"


def signature(data: dict) -> str:
    """Hash of the published content (not of the file bytes, which embed a
    timestamp) - equal signatures mean there is nothing new to upload."""
    return hashlib.sha256(json.dumps(data, sort_keys=True, default=str).encode()).hexdigest()


def _money(value: Any) -> Optional[str]:
    """Canonical 2dp string for comparing amounts from Excel and AEGIS."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return str(Decimal(str(value).replace(",", "").strip()).quantize(Decimal("0.01")))
    except Exception:  # noqa: BLE001 - not a number: the caller reports it
        return "invalid"


def _iso_date(value: Any) -> Optional[str]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text_value = str(value).strip()[:10]
    try:
        return datetime.strptime(text_value, "%Y-%m-%d").date().isoformat()
    except ValueError:
        return "invalid"


def _alloc_state(a: dict) -> dict:
    return {"line": str(a["line_id"]), "project": str(a["project_id"]) if a["project_id"] else None,
            "category": a["category"], "amount": _money(a["amount"]), "date": _iso_date(a["allocation_date"]),
            "description": _clean(a["description"])}


def snapshot(data: dict) -> dict:
    """What AEGIS puts in the editable cells - the baseline Excel edits are
    detected against: per bank line, and per split / cash-use part."""
    return {
        "lines": {
            str(l["id"]): {"category": l["category"], "project": str(l["project_id"]) if l["project_id"] else None,
                           "who": _clean(l["counterparty_name"]), "note": _clean(l["notes"])}
            for l in data["lines"]
        },
        "allocations": {str(a["id"]): _alloc_state(a) for a in data["allocations"]},
    }


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def _table_sheet(wb: Workbook, title: str, table_name: str, headers: list[str], rows: list[list[Any]],
                 widths: list[int], money_cols: set[int], date_cols: set[int],
                 editable_cols: Optional[set[int]] = None) -> Any:
    editable_cols = editable_cols or set()
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for i, cell in enumerate(ws[1], start=1):
        cell.fill = EDIT_HEADER_FILL if i in editable_cols else HEADER_FILL
        cell.font, cell.alignment = HEADER_FONT, Alignment(vertical="center")
    for col in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = MONEY_FORMAT
    for col in date_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = "yyyy-mm-dd"
    for col in editable_cols:
        for (cell,) in ws.iter_rows(min_row=2, max_row=max(2, len(rows) + 1), min_col=col, max_col=col):
            cell.protection = Protection(locked=False)
            cell.fill = EDIT_CELL_FILL
    ws.freeze_panes = "A2"
    last = f"{get_column_letter(len(headers))}{max(2, len(rows) + 1)}"
    if not rows:  # an Excel table needs at least one data row
        ws.append([None] * len(headers))
    table = Table(displayName=table_name, ref=f"A1:{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.protection.sheet = True
    ws.protection.autoFilter = False   # filtering/sorting still works on a protected sheet
    ws.protection.sort = False
    return ws


def build(data: dict, *, published_at: datetime) -> bytes:
    wb = Workbook()
    readme = wb.active
    readme.title = "Read Me"
    audit = data["audit"]
    todo = audit["to_do"]
    info = [
        ("AEGIS Bank & Project Money", None),
        ("Published by AEGIS", published_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("What this is", "The bank statement and where the money went, kept in step with AEGIS."),
        ("What you can edit", "On Bank Lines, the orange columns: Category, Project, Who and Note. Pick Category and Project "
                              "from the dropdowns. Everything else is locked."),
        ("What happens next", "Within about 2 minutes AEGIS applies your edit (claims, costs, petty cash and the ledger update "
                              "together) and fills in Sync status. Clear a cell to remove that tag."),
        ("If AEGIS changed it too", "AEGIS keeps its value and Sync status says Conflict - check it in AEGIS."),
        ("Splitting a line / cash uses", "On Splits & Cash Uses: change Project, Category, Amount, Date or What it paid for on "
                                         "a part, or clear its Amount to remove it. To add a part, use a blank row at the bottom: "
                                         "put the bank line's Line ID (or its first 8 characters, from Bank Lines) in Line, then "
                                         "the rest. On a cash withdrawal each part is what that cash paid for; anything not "
                                         "recorded stays in HQ Petty Cash. Parts can't add up to more than the line."),
        (None, None),
        ("Books check", "All checks pass" if audit["all_ok"] else "Something doesn't tie out - see below"),
    ]
    info += [(f"  {c['check']}", ("OK" if c["ok"] else "CHECK") + f" - {c['detail']}") for c in audit["checks"]]
    info += [
        (None, None),
        ("Still to do", None),
        ("  Lines with no category or project", int(todo["unclassified_lines"])),
        ("  Unidentified money in (in Suspense)", _num(todo["unclassified_money_in"])),
        ("  Unidentified money out (in Suspense)", _num(todo["unclassified_money_out"])),
        ("  Cash withdrawn, not yet accounted for", _num(todo["hq_petty_cash_not_yet_accounted_for"])),
        ("  Project costs with no cost type", _num(todo["unclassified_project_costs"])),
    ]
    if data.get("problems"):
        info += [(None, None), ("New parts from Excel that couldn't be applied (last 7 days) - please re-enter", None)]
        info += [(f"  {p['at']}  {p['new_value'] or ''}", p["message"]) for p in data["problems"]]
    for label, value in info:
        readme.append([label, value])
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 44
    readme.column_dimensions["B"].width = 110
    for row in readme.iter_rows(min_row=2):
        if isinstance(row[1].value, float):
            row[1].number_format = MONEY_FORMAT
    readme.protection.sheet = True

    lines_ws = _table_sheet(
        wb, "Bank Lines", "BankLines", BANK_LINE_HEADERS,
        [[str(l["id"]), l["transaction_date"], l["reference"], l["description"],
          _num(l["amount"]) if l["amount"] > 0 else None, _num(-l["amount"]) if l["amount"] < 0 else None,
          _num(l["balance"]), _label(l["category"]) or None, l["project"], l["counterparty_name"], l["notes"],
          int(l["parts"]) or None, l["journal_number"], l["bank_account"], l["sync_status"]] for l in data["lines"]],
        [38, 12, 20, 70, 14, 14, 14, 24, 34, 26, 30, 10, 20, 16, 44],
        money_cols={5, 6, 7}, date_cols={2}, editable_cols={8, 9, 10, 11},
    )
    split_rows = [
        [str(a["id"]), str(a["line_id"]), a["line_date"], a["reference"], _num(abs(a["line_amount"])),
         "Use of withdrawn cash" if a["line_category"] == "cash_withdrawal" else "Split of bank line",
         a["project"], _label(a["category"]) or None, _num(a["amount"]), a["allocation_date"], a["description"],
         a["sync_status"]]
        for a in data["allocations"]
    ] + [[None] * len(SPLIT_HEADERS) for _ in range(SPARE_SPLIT_ROWS)]
    splits_ws = _table_sheet(
        wb, "Splits & Cash Uses", "Allocations", SPLIT_HEADERS, split_rows,
        [38, 38, 12, 20, 14, 22, 34, 24, 14, 12, 50, 44], money_cols={5, 9}, date_cols={3, 10},
        editable_cols={7, 8, 9, 10, 11},
    )
    # In the spare rows the Line cell is editable too - that's how a new part names its bank line
    first_spare = len(data["allocations"]) + 2
    for (cell,) in splits_ws.iter_rows(min_row=first_spare, max_row=first_spare + SPARE_SPLIT_ROWS - 1, min_col=2, max_col=2):
        cell.protection = Protection(locked=False)
        cell.fill = EDIT_CELL_FILL
    for (cell,) in splits_ws.iter_rows(min_row=first_spare, max_row=first_spare + SPARE_SPLIT_ROWS - 1, min_col=10, max_col=10):
        cell.number_format = "yyyy-mm-dd"
    for (cell,) in splits_ws.iter_rows(min_row=first_spare, max_row=first_spare + SPARE_SPLIT_ROWS - 1, min_col=9, max_col=9):
        cell.number_format = MONEY_FORMAT
    _table_sheet(
        wb, "Projects", "Projects",
        ["Project", "Code", "Client", "Status", "Contract value", "Certified", "Collected", "Actual cost",
         "Margin to date", "Bank in", "Bank out", "Cash used"],
        [[p["name"], p["project_code"], p["client_name"], p["status"], _num(p["contract_value"]), _num(p["certified"]),
          _num(p["collected"]), _num(p["actual_cost"]), _num(Decimal(str(p["collected"])) - Decimal(str(p["actual_cost"]))),
          _num(p["bank_in"]), _num(p["bank_out"]), _num(p["cash_used"])] for p in data["projects"]],
        [38, 10, 26, 16, 16, 16, 16, 16, 16, 16, 16, 14], money_cols={5, 6, 7, 8, 9, 10, 11, 12}, date_cols=set(),
    )

    # Dropdown sources for the editable columns
    lists = wb.create_sheet("Lists")
    lists.append(["Categories", "Projects"])
    category_labels = [CATEGORY_LABELS[c] for c in EDITABLE_CATEGORIES]
    project_names = [p["name"] for p in data["projects"]]
    for i in range(max(len(category_labels), len(project_names))):
        lists.append([category_labels[i] if i < len(category_labels) else None,
                      project_names[i] if i < len(project_names) else None])
    lists.column_dimensions["A"].width = 30
    lists.column_dimensions["B"].width = 44
    lists.protection.sheet = True
    last_row = len(data["lines"]) + 1
    for column, source in (("H", f"Lists!$A$2:$A${len(category_labels) + 1}"),
                           ("I", f"Lists!$B$2:$B${max(2, len(project_names) + 1)}")):
        validation = DataValidation(type="list", formula1=f"={source}", allow_blank=True, showErrorMessage=True,
                                    errorTitle="Pick from the list", error="Choose a value from the dropdown, or clear the cell.")
        validation.add(f"{column}2:{column}{max(2, last_row)}")
        lines_ws.add_data_validation(validation)
    last_split_row = len(split_rows) + 1
    for column, source in (("H", f"Lists!$A$2:$A${len(category_labels) + 1}"),
                           ("G", f"Lists!$B$2:$B${max(2, len(project_names) + 1)}")):
        validation = DataValidation(type="list", formula1=f"={source}", allow_blank=True, showErrorMessage=True,
                                    errorTitle="Pick from the list", error="Choose a value from the dropdown, or clear the cell.")
        validation.add(f"{column}2:{column}{last_split_row}")
        splits_ws.add_data_validation(validation)
    amount_check = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                                  showErrorMessage=True, errorTitle="Amount", error="Enter an amount above zero, or clear the cell.")
    amount_check.add(f"I2:I{last_split_row}")
    splits_ws.add_data_validation(amount_check)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Reading edits back
# ---------------------------------------------------------------------------

def read_edits(content: bytes) -> dict[str, dict[str, Optional[str]]]:
    """{line_id: {field: raw cell text}} for the editable columns of Bank Lines."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb["Bank Lines"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    try:
        id_col = header.index("Line ID")
    except ValueError:
        return {}
    cols = {field: header.index(name) for name, field in EDITABLE_HEADERS.items() if name in header}
    out: dict[str, dict[str, Optional[str]]] = {}
    for row in rows:
        line_id = _clean(row[id_col]) if id_col < len(row) else None
        if not line_id:
            continue
        out[line_id] = {field: _clean(row[i]) if i < len(row) else None for field, i in cols.items()}
    wb.close()
    return out


async def _apply_edits(db: AsyncSession, *, org_id: str, edits: dict[str, dict], baseline: dict[str, dict],
                       editor_name: Optional[str], editor_user_id: Optional[UUID]) -> dict:
    """Compare edited cells with the published baseline and AEGIS's current
    values; apply genuine edits, log every outcome. Returns counts."""
    counts = {"applied": 0, "rejected": 0, "conflict": 0}
    if not edits or not baseline:
        return counts
    label_to_code = {v.lower(): k for k, v in CATEGORY_LABELS.items() if k in EDITABLE_CATEGORIES}
    label_to_code.update({k: k for k in EDITABLE_CATEGORIES})
    projects: dict[str, list] = {}
    deleted_projects: dict[str, str] = {}
    for r in await db.execute(text("SELECT id, name, is_deleted FROM projects.projects WHERE organization_id = :o"),
                              {"o": org_id}):
        if r.is_deleted:
            deleted_projects.setdefault(r.name.strip().lower(), str(r.id))
        else:
            projects.setdefault(r.name.strip().lower(), []).append(r.id)
    current = {
        str(r["id"]): {"category": r["category"], "project": str(r["project_id"]) if r["project_id"] else None,
                       "who": _clean(r["counterparty_name"]), "note": _clean(r["notes"]), "project_name": r["project_name"]}
        for r in (await db.execute(text("""
            SELECT l.id, l.category, l.project_id, l.counterparty_name, l.notes, p.name AS project_name
            FROM finance.bank_statement_lines l LEFT JOIN projects.projects p ON p.id = l.project_id
            WHERE l.organization_id = :o
        """), {"o": org_id})).mappings()
    }

    def normalise(field: str, raw: Optional[str]) -> tuple[Optional[str], Optional[str]]:
        """-> (stored value, error)"""
        if raw is None:
            return None, None
        if field == "category":
            code = label_to_code.get(raw.lower())
            return (code, None) if code else (None, f"'{raw}' is not a category")
        if field == "project":
            ids = projects.get(raw.lower(), [])
            if len(ids) == 1:
                return str(ids[0]), None
            if not ids and raw.lower() in deleted_projects:
                # Still shown on lines tagged before the project was deleted:
                # fine if untouched, refused if someone picks it anew.
                return deleted_projects[raw.lower()], None
            return None, f"project '{raw}' {'is ambiguous' if ids else 'was not found'}"
        limit = FIELD_LIMITS[field]
        return (raw[:limit], None) if len(raw) <= limit else (None, f"{field} is longer than {limit} characters")

    log: list[dict] = []
    updates_by_line: dict[str, dict] = {}
    for line_id, cells in edits.items():
        base, now = baseline.get(line_id), current.get(line_id)
        if base is None or now is None:
            continue   # a line published before / not in AEGIS any more
        for field, raw in cells.items():
            value, error = normalise(field, raw)
            if error is None and value == base.get(field):
                continue                      # nobody touched it
            if error is None and value == now.get(field):
                continue                      # already what AEGIS has
            if field == "project" and value in deleted_projects.values():
                error = f"project '{raw}' has been deleted in AEGIS"
            shown_old = now["project_name"] if field == "project" else now.get(field)
            entry = {"line_id": line_id, "field": field, "old_value": shown_old, "new_value": raw}
            if error:
                if raw == (_label(base.get(field)) if field == "category" else base.get(field)):
                    continue
                entry.update(outcome="rejected", message=error)
            elif now.get(field) != base.get(field):
                entry.update(outcome="conflict",
                             message=f"{field} was also changed in AEGIS - kept '{shown_old or ''}', your value '{raw or ''}' was not applied")
            else:
                entry.update(outcome="applied", message=None)
                updates_by_line.setdefault(line_id, {})[FIELD_TO_COLUMN[field]] = value
            counts[entry["outcome"]] += 1
            log.append(entry)

    user_id = str(editor_user_id) if editor_user_id else None
    for line_id, updates in updates_by_line.items():
        await reconciliation.tag_lines(db, org_id=org_id, user_id=user_id, updates=updates, line_ids=[UUID(line_id)])
    if updates_by_line:
        await bank_books.sync(db, org_id=org_id, user_id=user_id, line_ids=[UUID(i) for i in updates_by_line])
    if log:
        await db.execute(
            text("""
                INSERT INTO finance.bank_workbook_changes (
                    organization_id, line_id, field, old_value, new_value, outcome, message, edited_by_name, edited_by_user_id
                )
                SELECT :org_id, x.line_id, x.field, x.old_value, x.new_value, x.outcome, x.message, :name, :user_id
                FROM jsonb_to_recordset(CAST(:log AS jsonb)) AS x(
                    line_id uuid, field text, old_value text, new_value text, outcome text, message text)
            """),
            {"org_id": org_id, "name": editor_name, "user_id": editor_user_id, "log": json.dumps(log)},
        )
    return counts


def read_split_edits(content: bytes) -> list[dict]:
    """Every row of Splits & Cash Uses that has anything in it:
    {"id": allocation id or None (a spare row), "line": raw Line cell, field: raw cell}."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if "Splits & Cash Uses" not in wb.sheetnames:
        wb.close()
        return []
    rows = wb["Splits & Cash Uses"].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    if "Allocation ID" not in header or "Line" not in header:
        wb.close()
        return []
    id_col, line_col = header.index("Allocation ID"), header.index("Line")
    cols = {field: header.index(name) for name, field in SPLIT_EDITABLE.items() if name in header}
    out = []
    for row in rows:
        cell = lambda i: row[i] if i < len(row) else None  # noqa: E731
        entry = {"id": _clean(cell(id_col)), "line": _clean(cell(line_col))}
        entry.update({field: cell(i) for field, i in cols.items()})
        if entry["id"] or any(v not in (None, "") for k, v in entry.items() if k != "id"):
            out.append(entry)
    wb.close()
    return out


async def _project_lookup(db: AsyncSession, org_id: str) -> tuple[dict, dict, dict]:
    """(live name -> [ids], deleted name -> id, id -> name)"""
    live: dict[str, list] = {}
    deleted: dict[str, str] = {}
    names: dict[str, str] = {}
    for r in await db.execute(text("SELECT id, name, is_deleted FROM projects.projects WHERE organization_id = :o"), {"o": org_id}):
        names[str(r.id)] = r.name
        if r.is_deleted:
            deleted.setdefault(r.name.strip().lower(), str(r.id))
        else:
            live.setdefault(r.name.strip().lower(), []).append(r.id)
    return live, deleted, names


async def _apply_split_edits(db: AsyncSession, *, org_id: str, rows: list[dict], baseline: dict[str, dict],
                             applied_keys: set[str], editor_name: Optional[str],
                             editor_user_id: Optional[UUID]) -> tuple[dict, list[str]]:
    """Splits & Cash Uses edits: change / clear existing parts, add new ones.
    Each affected line's parts are saved together through
    bank_books.replace_allocations (same rules as the AEGIS split editor),
    in a savepoint so one bad line doesn't stop the others. Returns
    (counts, keys of new rows applied this cycle)."""
    counts = {"applied": 0, "rejected": 0, "conflict": 0}
    new_keys: list[str] = []
    if not rows:
        return counts, new_keys
    live, deleted, names = await _project_lookup(db, org_id)
    label_to_code = {v.lower(): k for k, v in CATEGORY_LABELS.items()}
    label_to_code.update({k: k for k in CATEGORY_LABELS})
    lines = {str(r.id): r for r in await db.execute(text("""
        SELECT id, amount, category, transaction_date FROM finance.bank_statement_lines WHERE organization_id = :o
    """), {"o": org_id})}
    current = {str(r["id"]): _alloc_state(r) for r in (await db.execute(text("""
        SELECT id, line_id, project_id, category, amount, allocation_date, description
        FROM finance.bank_line_allocations WHERE organization_id = :o
    """), {"o": org_id})).mappings()}

    def resolve_line(raw: Optional[str]) -> tuple[Optional[str], Optional[str]]:
        if not raw:
            return None, "Line is empty - paste the bank line's Line ID (or its first 8 characters)"
        key = raw.strip().lower()
        if key in lines:
            return key, None
        if len(key) >= SHORT_ID_MIN:
            matches = [i for i in lines if i.startswith(key)]
            if len(matches) == 1:
                return matches[0], None
            if len(matches) > 1:
                return None, f"'{raw}' matches more than one bank line - use more of the Line ID"
        return None, f"no bank line with ID '{raw}'"

    def normalise(row: dict, line_id: str) -> tuple[dict, Optional[str]]:
        errors = []
        project = _clean(row.get("project"))
        project_id = None
        if project:
            ids = live.get(project.lower(), [])
            if len(ids) == 1:
                project_id = str(ids[0])
            elif not ids and project.lower() in deleted:
                project_id = deleted[project.lower()]
            else:
                errors.append(f"project '{project}' {'is ambiguous' if ids else 'was not found'}")
        category = _clean(row.get("category"))
        code = None
        if category:
            code = label_to_code.get(category.lower())
            if not code:
                errors.append(f"'{category}' is not a category")
        amount = _money(row.get("amount"))
        if amount == "invalid":
            errors.append("Amount is not a number")
        date = _iso_date(row.get("date"))
        if date == "invalid":
            errors.append("Date should look like 2026-02-08")
        if not date or date == "invalid":
            date = lines[line_id].transaction_date.isoformat()
        description = _clean(row.get("description"))
        state = {"line": line_id, "project": project_id, "category": code,
                 "amount": None if amount == "invalid" else amount, "date": date,
                 "description": description[:2000] if description else None}
        return state, "; ".join(errors) or None

    def same(a: Optional[dict], b: Optional[dict]) -> bool:
        keys = ("project", "category", "amount", "date", "description")
        return a is not None and b is not None and all(a.get(k) == b.get(k) for k in keys)

    def describe(state: Optional[dict]) -> Optional[str]:
        if not state:
            return None
        parts = [names.get(state.get("project") or "", "no project"), _label(state.get("category")) or "no category",
                 state.get("amount") or "no amount", state.get("date") or "", state.get("description") or ""]
        return " · ".join(p for p in parts if p)

    log: list[dict] = []
    per_line: dict[str, dict] = {}
    for row in rows:
        aid = row["id"]
        if aid:
            base = baseline.get(aid)
            if base is None:
                continue
            excel, error = normalise(row, base["line"])
            deleting = excel["amount"] is None and not error
            if not error and not deleting and same(excel, base):
                continue                                    # nobody touched it
            cur = current.get(aid)
            if deleting and cur is None:
                continue                                    # already removed
            if not error and not deleting and same(excel, cur):
                continue                                    # already what AEGIS has
            entry = {"line_id": base["line"], "allocation_id": aid, "field": "split",
                     "old_value": describe(cur or base), "new_value": "(removed)" if deleting else describe(excel)}
            if error:
                entry.update(outcome="rejected", message=error)
            elif cur is None:
                entry.update(outcome="conflict", message="this part was removed in AEGIS - your change was not applied")
            elif not same(cur, base):
                entry.update(outcome="conflict", message="this part was also changed in AEGIS - AEGIS's version was kept")
            else:
                bucket = per_line.setdefault(base["line"], {"updates": {}, "new": []})
                bucket["updates"][aid] = None if deleting else excel
                bucket.setdefault("entries", []).append(entry)
                continue
            counts[entry["outcome"]] += 1
            log.append(entry)
        else:
            key = hashlib.sha256(json.dumps(
                [str(row.get(k) or "") for k in ("line", "project", "category", "amount", "date", "description")]
            ).encode()).hexdigest()
            if key in applied_keys:
                continue                                    # applied last cycle; republish hasn't landed yet
            line_id, error = resolve_line(row.get("line"))
            excel = None
            if line_id:
                excel, error = normalise(row, line_id)
                if not error and excel["amount"] is None:
                    error = "enter an Amount"
            shown = describe(excel) if excel else " · ".join(str(row.get(k)) for k in ("line", "project", "category", "amount") if row.get(k))
            entry = {"line_id": line_id, "allocation_id": None, "field": "split", "old_value": None, "new_value": shown}
            if error:
                entry.update(outcome="rejected", message=error)
                counts["rejected"] += 1
                log.append(entry)
                continue
            bucket = per_line.setdefault(line_id, {"updates": {}, "new": []})
            bucket["new"].append((excel, entry, key))

    user_id = str(editor_user_id) if editor_user_id else None
    for line_id, bucket in per_line.items():
        line = lines[line_id]
        is_cash = line.category == "cash_withdrawal" and line.amount < 0
        final = []
        for aid, cur in current.items():
            if cur["line"] != line_id:
                continue
            state = bucket["updates"].get(aid, cur) if aid in bucket["updates"] else cur
            if state is None:
                continue                                    # cleared -> removed
            final.append({"id": UUID(aid), **state})
        final += [dict(state) for state, _, _ in bucket["new"]]
        entries = bucket.get("entries", []) + [entry for _, entry, _ in bucket["new"]]
        error = None
        if is_cash and any((p.get("category") or "") in NOT_A_CASH_USE for p in final):
            error = "a use of withdrawn cash can't be categorised as " + ", ".join(
                sorted({_label(p["category"]) for p in final if (p.get("category") or "") in NOT_A_CASH_USE}))
        if not error:
            try:
                before = {aid for aid, cur in current.items() if cur["line"] == line_id}
                async with db.begin_nested():
                    await bank_books.replace_allocations(db, org_id=org_id, user_id=user_id, line_id=UUID(line_id), allocations=[
                        {"id": p.get("id"), "project_id": UUID(p["project"]) if p.get("project") else None,
                         "category": p.get("category"), "amount": Decimal(p["amount"]), "description": p.get("description"),
                         "allocation_date": date.fromisoformat(p["date"]) if p.get("date") else None}
                        for p in final
                    ])
                created = [str(r.id) for r in await db.execute(text("""
                    SELECT id FROM finance.bank_line_allocations WHERE line_id = :l ORDER BY created_at, id
                """), {"l": line_id}) if str(r.id) not in before]
                for (_, entry, key), new_id in zip(bucket["new"], created):
                    entry["allocation_id"] = new_id
                    new_keys.append(key)
            except GeneralLedgerError as exc:
                error = exc.detail
        for entry in entries:
            entry.update(outcome="rejected", message=error) if error else entry.update(outcome="applied", message=None)
            counts[entry["outcome"]] += 1
            log.append(entry)

    if log:
        await db.execute(
            text("""
                INSERT INTO finance.bank_workbook_changes (
                    organization_id, line_id, allocation_id, field, old_value, new_value, outcome, message,
                    edited_by_name, edited_by_user_id
                )
                SELECT :org_id, x.line_id, x.allocation_id, x.field, x.old_value, x.new_value, x.outcome, x.message, :name, :user_id
                FROM jsonb_to_recordset(CAST(:log AS jsonb)) AS x(
                    line_id uuid, allocation_id uuid, field text, old_value text, new_value text, outcome text, message text)
            """),
            {"org_id": org_id, "name": editor_name, "user_id": editor_user_id, "log": json.dumps(log)},
        )
    return counts, new_keys


# ---------------------------------------------------------------------------
# Publication row, lease
# ---------------------------------------------------------------------------

async def status(db: AsyncSession, org_id: str) -> Optional[dict]:
    row = (await db.execute(
        text("""
            SELECT organization_id, file_name, drive_id, item_id, web_url, last_published_at, last_attempt_at,
                   last_read_at, last_status, last_error, rows_published
            FROM finance.bank_workbook_publications WHERE organization_id = :org_id
        """),
        {"org_id": org_id},
    )).mappings().first()
    return dict(row) if row else None


async def recent_changes(db: AsyncSession, org_id: str, limit: int = 20) -> list[dict]:
    rows = await db.execute(
        text("""
            SELECT c.created_at, c.field, c.old_value, c.new_value, c.outcome, c.message, c.edited_by_name,
                   l.transaction_date, l.reference, l.amount
            FROM finance.bank_workbook_changes c LEFT JOIN finance.bank_statement_lines l ON l.id = c.line_id
            WHERE c.organization_id = :org_id ORDER BY c.created_at DESC LIMIT :limit
        """),
        {"org_id": org_id, "limit": limit},
    )
    return [dict(r) for r in rows.mappings()]


async def _record(db: AsyncSession, org_id: str, **values: Any) -> None:
    json_columns = ("published_snapshot", "applied_new_row_keys")
    for k in json_columns:
        if k in values:
            values[k] = json.dumps(values[k])
    assignments = ", ".join(
        f"{k} = CAST(:{k} AS jsonb)" if k in json_columns else f"{k} = :{k}" for k in values
    )
    await db.execute(
        text(f"UPDATE finance.bank_workbook_publications SET {assignments}, updated_at = NOW() WHERE organization_id = :org_id"),
        {"org_id": org_id, **values},
    )
    await db.commit()


async def _take_lease(db: AsyncSession, org_id: str) -> Optional[dict]:
    await db.execute(
        text("""
            INSERT INTO finance.bank_workbook_publications (organization_id, file_name)
            VALUES (:org_id, :file_name) ON CONFLICT (organization_id) DO NOTHING
        """),
        {"org_id": org_id, "file_name": WORKBOOK_FILE_NAME},
    )
    row = (await db.execute(
        text(f"""
            UPDATE finance.bank_workbook_publications
            SET sync_lease_until = NOW() + INTERVAL '{LEASE_MINUTES} minutes'
            WHERE organization_id = :org_id AND (sync_lease_until IS NULL OR sync_lease_until < NOW())
            RETURNING drive_id, item_id, published_etag, published_snapshot, content_signature, applied_new_row_keys
        """),
        {"org_id": org_id},
    )).mappings().first()
    await db.commit()
    if not row:
        return None
    lease = dict(row)
    for k in ("published_snapshot", "applied_new_row_keys"):
        if isinstance(lease[k], str):
            lease[k] = json.loads(lease[k])
    return lease


async def _release_lease(db: AsyncSession, org_id: str) -> None:
    await db.execute(text("UPDATE finance.bank_workbook_publications SET sync_lease_until = NULL WHERE organization_id = :org_id"),
                     {"org_id": org_id})
    await db.commit()


async def _editor(db: AsyncSession, item: dict) -> tuple[Optional[str], Optional[UUID]]:
    user = ((item.get("lastModifiedBy") or {}).get("user")) or {}
    name, email = user.get("displayName"), user.get("email")
    user_id = None
    if email:
        user_id = (await db.execute(
            text("SELECT id FROM core.users WHERE lower(email) = lower(:email) LIMIT 1"), {"email": email},
        )).scalar()
    return (f"{name} <{email}>" if name and email else name or email), user_id


# ---------------------------------------------------------------------------
# The sync cycle
# ---------------------------------------------------------------------------

async def sync(db: AsyncSession, *, org_id: str, force: bool = False) -> Optional[dict]:
    """Read Excel edits back (if the file changed), then republish. Returns
    the publication status row, or None if another cycle holds the lease.
    Never raises for Microsoft-side failures - they are recorded on the row."""
    lease = await _take_lease(db, org_id)
    if lease is None:
        return None
    now = datetime.now(timezone.utc)
    try:
        connection = await data_room_sync._load_connection(db, UUID(org_id))
        client = GraphClient(tenant_id=connection.tenant_id)
        read_etag: Optional[str] = None
        edits_applied = False

        # 1-2. Read edits back when the file changed since AEGIS last wrote it
        if lease["item_id"]:
            try:
                item = await client.get(f"drives/{lease['drive_id']}/items/{lease['item_id']}")
            except GraphNotFoundError:
                item = None   # deleted in SharePoint - publish a fresh copy below
                lease["item_id"] = None
            if item:
                read_etag = item.get("eTag")
                if read_etag and read_etag != lease["published_etag"]:
                    content = await data_room_sync.fetch_bytes(
                        db, organization_id=UUID(org_id), drive_id=lease["drive_id"], item_id=lease["item_id"])
                    editor_name, editor_user_id = await _editor(db, item)
                    snap = lease["published_snapshot"] or {}
                    # snapshots from before editable splits were just {line_id: {...}}
                    line_base = snap.get("lines", {}) if "lines" in snap else snap
                    split_base = snap.get("allocations", {})
                    counts = await _apply_edits(
                        db, org_id=org_id, edits=read_edits(content), baseline=line_base,
                        editor_name=editor_name, editor_user_id=editor_user_id)
                    split_counts, new_keys = await _apply_split_edits(
                        db, org_id=org_id, rows=read_split_edits(content), baseline=split_base,
                        applied_keys=set(lease["applied_new_row_keys"] or []),
                        editor_name=editor_name, editor_user_id=editor_user_id)
                    await db.commit()
                    edits_applied = any(counts.values()) or any(split_counts.values())
                    logger.info("bank_workbook.edits_read", organization_id=org_id, lines=counts, splits=split_counts)
                    values = {"last_read_at": now}
                    if new_keys:
                        values["applied_new_row_keys"] = list(lease["applied_new_row_keys"] or []) + new_keys
                    await _record(db, org_id, **values)

        # 3. Republish (skip when nothing changed on either side)
        data = await collect(db, org_id)
        sig = signature(data)
        if not force and not edits_applied and lease["item_id"] and sig == lease["content_signature"]:
            # Nothing to write. If the eTag moved with no real edit (Excel
            # re-saved the file), remember it so we don't re-read every cycle.
            values = {"last_attempt_at": now, "last_status": "unchanged", "last_error": None}
            if read_etag:
                values["published_etag"] = read_etag
            await _record(db, org_id, **values)
            return await status(db, org_id)

        content = build(data, published_at=now)
        headers = {"If-Match": read_etag} if (read_etag and lease["item_id"]) else None
        try:
            body = await client.put_content(
                f"drives/{connection.drive_id}/items/{connection.root_item_id}:/{quote(WORKBOOK_FILE_NAME)}:/content",
                content=content, content_type=XLSX_MIME, extra_headers=headers,
            ) if len(content) <= sharepoint.SIMPLE_UPLOAD_MAX_BYTES else None
            if body is None:   # very large workbook: resumable upload (no If-Match support there)
                item = await sharepoint.upload_large_file(client, connection.drive_id, connection.root_item_id,
                                                          WORKBOOK_FILE_NAME, content, XLSX_MIME)
                body = {"id": item.item_id, "webUrl": item.web_url, "eTag": item.etag}
        except GraphConflictError:
            # Someone edited while we were working - don't overwrite them.
            await _record(db, org_id, last_attempt_at=now, last_status="retry",
                          last_error="The workbook was edited while AEGIS was updating it - will retry next cycle.")
            return await status(db, org_id)

        await _record(
            db, org_id, drive_id=connection.drive_id, item_id=body["id"], web_url=body.get("webUrl"),
            published_etag=body.get("eTag"), published_snapshot=snapshot(data), content_signature=sig,
            last_published_at=now, last_attempt_at=now, last_status="published", last_error=None,
            rows_published=len(data["lines"]), applied_new_row_keys=[],
        )
        return await status(db, org_id)
    except Exception as exc:  # noqa: BLE001 - recorded on the row, never breaks the caller
        await db.rollback()
        logger.error("bank_workbook.sync_failed", organization_id=org_id, error=str(exc))
        await _record(db, org_id, last_attempt_at=now, last_status="failed", last_error=str(exc)[:2000])
        return await status(db, org_id)
    finally:
        await _release_lease(db, org_id)


# Phase 1 name, kept for callers
async def publish(db: AsyncSession, *, org_id: str, force: bool = False) -> Optional[dict]:
    return await sync(db, org_id=org_id, force=force)
