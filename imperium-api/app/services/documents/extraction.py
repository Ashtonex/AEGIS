"""Reads text back out of an uploaded PDF/Word document.

Unlike renderers.py (which only ever GENERATES documents AEGIS produces),
this is the other direction: given the raw bytes of a file a user uploaded
anywhere in the system, get its plain text content so it can be searched and
referenced by the AI assistant.

Does not implement interfaces.TextExtractor - that ABC takes a local
file_path, but extraction here always starts from bytes already fetched from
remote storage (Supabase Storage or SharePoint), so writing them to a
temporary file first would be pure overhead with no benefit.

Scanned/image-only PDFs (no real text layer) are explicitly out of scope -
pdfplumber returning no text for every page is recorded as NO_TEXT_FOUND, a
normal outcome, not a failure. OCR is a separate, heavier addition.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from enum import Enum
from pathlib import PurePosixPath
from typing import Optional

import docx
import pdfplumber

from core.config import settings


class ExtractionStatus(str, Enum):
    EXTRACTED = "extracted"
    NO_TEXT_FOUND = "no_text_found"
    UNSUPPORTED_FORMAT = "unsupported_format"
    FAILED = "failed"


@dataclass
class ExtractionResult:
    status: ExtractionStatus
    text: Optional[str] = None
    error: Optional[str] = None


_PDF_EXTENSIONS = {".pdf"}
_DOCX_EXTENSIONS = {".docx"}
# Legacy binary Word format - python-docx cannot open these at all. Flagged
# explicitly so a .doc upload is reported as "we can't read this format" and
# not silently misclassified as "we tried and found no text" (NO_TEXT_FOUND).
_LEGACY_DOC_EXTENSIONS = {".doc"}

_PDF_MIME_TYPES = {"application/pdf"}
_DOCX_MIME_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
_LEGACY_DOC_MIME_TYPES = {"application/msword"}


def _resolve_extension(file_name: Optional[str], mime_type: Optional[str]) -> Optional[str]:
    if file_name:
        ext = PurePosixPath(file_name).suffix.lower()
        if ext:
            return ext
    if mime_type in _PDF_MIME_TYPES:
        return ".pdf"
    if mime_type in _DOCX_MIME_TYPES:
        return ".docx"
    if mime_type in _LEGACY_DOC_MIME_TYPES:
        return ".doc"
    return None


def _extract_pdf(content: bytes) -> ExtractionResult:
    pages_text: list[str] = []
    max_pages = settings.DOCUMENT_EXTRACTION_MAX_PAGES
    truncated = False
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for index, page in enumerate(pdf.pages):
            if index >= max_pages:
                truncated = True
                break
            pages_text.append(page.extract_text() or "")

    text = "\n".join(pages_text).strip()
    if not text:
        # Most commonly a scanned/image-only PDF with no real text layer -
        # out of scope for this pass (would need OCR), so this is a normal,
        # expected outcome rather than an error.
        return ExtractionResult(status=ExtractionStatus.NO_TEXT_FOUND)
    if truncated:
        text += f"\n\n[Extraction truncated after {max_pages} pages.]"
    return ExtractionResult(status=ExtractionStatus.EXTRACTED, text=text)


def _extract_docx(content: bytes) -> ExtractionResult:
    document = docx.Document(io.BytesIO(content))
    parts: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    text = "\n".join(parts).strip()
    if not text:
        return ExtractionResult(status=ExtractionStatus.NO_TEXT_FOUND)
    return ExtractionResult(status=ExtractionStatus.EXTRACTED, text=text)


def extract_text(content: bytes, *, file_name: Optional[str] = None, mime_type: Optional[str] = None) -> ExtractionResult:
    """Pure, synchronous, no I/O beyond the bytes already given - callers
    (the arq job) are responsible for fetching content and are expected to
    run this via asyncio.to_thread, since it's CPU-bound and would otherwise
    block the worker's event loop for the duration of a large PDF parse."""
    extension = _resolve_extension(file_name, mime_type)

    try:
        if extension in _PDF_EXTENSIONS:
            return _extract_pdf(content)
        if extension in _DOCX_EXTENSIONS:
            return _extract_docx(content)
        if extension in _LEGACY_DOC_EXTENSIONS:
            return ExtractionResult(
                status=ExtractionStatus.UNSUPPORTED_FORMAT,
                error="Legacy .doc is not supported - convert to .docx and re-upload.",
            )
        return ExtractionResult(
            status=ExtractionStatus.UNSUPPORTED_FORMAT,
            error=f"No text extractor for this file type (file_name={file_name!r}, mime_type={mime_type!r}).",
        )
    except Exception as exc:  # noqa: BLE001 - any parser failure is reported, never raised
        return ExtractionResult(status=ExtractionStatus.FAILED, error=str(exc))
